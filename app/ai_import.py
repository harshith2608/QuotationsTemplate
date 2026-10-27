"""Bring your old quotation: a firm uploads one quotation it has sent before (PDF or Excel) and its app is set up from
it: its items and rates (saved lines), a starting quotation (template), its payment terms and notes, and the quotation
itself rebuilt in QuoteRoom.

    1. upload (admins)   the text is taken out of the file here: PDF text (a scan has none, so it's refused) or the
                         Excel sheets' cells. Only that text goes to the AI; the file itself is never stored.
    2. read              Claude (Anthropic API, ANTHROPIC_API_KEY) returns its sections and lines as structured data.
                         Numbers are copied from the file, never estimated; the customer's details are left out.
    3. review            every line is editable; the page warns when the lines don't add up to the file's total.
    4. set up            the reviewed lines go through the same loader as the console's "Load price list"
                         (app/price_list.py), into this business only.

One switch turns it on or off for every trial and customer in the shared app: the console's AI import card
(public.console_state 'ai_import' in the trials database, read here every SWITCH_TTL seconds). A single-business
deployment turns it on with AI_IMPORT=1. Either way it also needs ANTHROPIC_API_KEY, and it's off in the demo.
Each business may read LIMIT_TRIAL files during its trial and LIMIT_MONTH a month on a paid plan; every read is
logged in public.ai_imports (trials database) for the console.
"""
import io
import json
import os
import re
import time
from datetime import datetime
from typing import List, Optional

from pydantic import BaseModel
from sqlalchemy import text

from . import db

MODEL = 'claude-haiku-5-5'
STATE_KEY = 'ai_import'
USAGE_KEY = 'ai_import_usage'
SWITCH_TTL = 30
LIMIT_TRIAL = 5
LIMIT_MONTH = 20
MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 60000
MIN_TEXT = 200
TIMEOUT = 50                    # seconds (the app's requests may run for 60)
LOG_DDL = """CREATE TABLE IF NOT EXISTS public.ai_imports (
    id         SERIAL PRIMARY KEY,
    slug       VARCHAR(40),
    at         TIMESTAMP NOT NULL DEFAULT (now() AT TIME ZONE 'utc'),
    kind       VARCHAR(10),
    ok         BOOLEAN NOT NULL,
    lines      INTEGER,
    tokens_in  INTEGER,
    tokens_out INTEGER,
    note       VARCHAR(200))"""

_switch = {}


class ImportProblem(ValueError):
    """A message for the person importing (the file can't be read, the limit is reached…)."""


# ── On / off ─────────────────────────────────────────────────────────────────

def _truthy(value):
    return (value or '').strip().lower() in ('1', 'true', 'yes', 'on')


def switched_on():
    """The platform switch: the console's, for the shared app; AI_IMPORT=1 for a single-business deployment."""
    from . import tenancy
    if not tenancy.tenant_mode():
        return _truthy(os.environ.get('AI_IMPORT'))
    hit = _switch.get('on')
    if hit and time.time() - hit[1] < SWITCH_TTL:
        return hit[0]
    on = False
    try:
        with db.engine.connect() as conn:
            if conn.execute(text("SELECT to_regclass('public.console_state')")).scalar() is not None:
                on = conn.execute(text('SELECT value FROM public.console_state WHERE key = :k'),
                                  {'k': STATE_KEY}).scalar() == '1'
    except Exception:
        on = False
    _switch['on'] = (on, time.time())
    return on


def set_switch(on):
    """The console's switch (the console's database is the trials database). Takes effect within SWITCH_TTL."""
    with db.engine.begin() as conn:
        conn.execute(text('DELETE FROM public.console_state WHERE key = :k'), {'k': STATE_KEY})
        conn.execute(text('INSERT INTO public.console_state (key, value) VALUES (:k, :v)'),
                     {'k': STATE_KEY, 'v': '1' if on else '0'})
    _switch.clear()


def console_switch():
    """The console's own reading of the switch (no cache)."""
    with db.engine.connect() as conn:
        return conn.execute(text('SELECT value FROM public.console_state WHERE key = :k'), {'k': STATE_KEY}).scalar() == '1'


def api_key():
    return (os.environ.get('ANTHROPIC_API_KEY') or '').strip()


def available():
    """Shown and usable in this app: switched on, the API key set, not the public demo."""
    from .demo import demo_enabled
    return not demo_enabled() and bool(api_key()) and switched_on()


# ── Limits ───────────────────────────────────────────────────────────────────

def _usage():
    from .models import AppSetting
    try:
        return json.loads(AppSetting.get(USAGE_KEY) or '{}')
    except ValueError:
        return {}


def left():
    """(files this business can still read, the limit's wording)."""
    from . import plan as plan_mod
    use = _usage()
    status = plan_mod.load_status()
    if status.has_limit and status.is_trial:
        return max(LIMIT_TRIAL - use.get('trial', 0), 0), f'{LIMIT_TRIAL} during your free trial'
    return max(LIMIT_MONTH - use.get(datetime.utcnow().strftime('%Y-%m'), 0), 0), f'{LIMIT_MONTH} a month'


def _count():
    from . import plan as plan_mod
    from .models import AppSetting
    use = _usage()
    status = plan_mod.load_status()
    key = 'trial' if status.has_limit and status.is_trial else datetime.utcnow().strftime('%Y-%m')
    use[key] = use.get(key, 0) + 1
    AppSetting.set(USAGE_KEY, json.dumps(use))
    db.session.commit()


def _log(kind, ok, lines=None, tokens_in=None, tokens_out=None, note=''):
    """One row per read in public.ai_imports (shared app only; never breaks the import)."""
    from flask import g
    from . import tenancy
    if not tenancy.tenant_mode():
        return
    try:
        with db.engine.begin() as conn:
            conn.execute(text(LOG_DDL))
            conn.execute(text('INSERT INTO public.ai_imports (slug, kind, ok, lines, tokens_in, tokens_out, note) '
                              'VALUES (:s, :k, :ok, :l, :ti, :to, :n)'),
                         {'s': g.get('tenant_slug'), 'k': kind, 'ok': ok, 'l': lines, 'ti': tokens_in,
                          'to': tokens_out, 'n': (note or '')[:200]})
    except Exception:
        pass


def recent(limit=30):
    """The latest reads, for the console (none before the first)."""
    with db.engine.connect() as conn:
        if conn.execute(text("SELECT to_regclass('public.ai_imports')")).scalar() is None:
            return []
        return [dict(r) for r in conn.execute(text('SELECT * FROM public.ai_imports ORDER BY id DESC LIMIT :n'),
                                              {'n': limit}).mappings()]


# ── 1. The file's text ───────────────────────────────────────────────────────

def file_text(filename, raw):
    """(kind, text) from an uploaded PDF or .xlsx. Raises ImportProblem for anything it can't read."""
    name = (filename or '').lower()
    if not raw:
        raise ImportProblem('Choose your quotation file (PDF or Excel).')
    if len(raw) > MAX_BYTES:
        raise ImportProblem('That file is larger than 5 MB. Upload the quotation on its own.')
    if name.endswith('.pdf') or raw[:5] == b'%PDF-':
        return 'pdf', _pdf_text(raw)
    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        return 'excel', _excel_text(raw)
    if name.endswith('.xls'):
        raise ImportProblem('This is an old Excel file (.xls). Open it in Excel, choose Save As → Excel Workbook '
                            '(.xlsx), and upload that.')
    raise ImportProblem('Upload a PDF or an Excel file (.xlsx).')


def _pdf_text(raw):
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ImportProblem('This PDF is password-protected. Upload a copy without a password.')
        body = '\n'.join((page.extract_text() or '') for page in reader.pages[:20])
    except ImportProblem:
        raise
    except Exception:
        raise ImportProblem("This PDF couldn't be opened. Try saving it again, or upload the Excel file instead.")
    if len(body.strip()) < MIN_TEXT:
        raise ImportProblem('This PDF is a scan (a picture of the pages), so there\'s no text to read. Upload the '
                            'original file you made it from: the PDF saved from Word or Excel, or the Excel file.')
    return body[:MAX_TEXT]


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float):
        return f'{value:.2f}'.rstrip('0').rstrip('.')
    return ' '.join(str(value).split())


def _excel_text(raw):
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ImportProblem("This Excel file couldn't be opened. Save it again as .xlsx and upload that.")
    out = []
    for sheet in book.worksheets[:5]:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append('\t'.join(cells))
        if rows:
            out.append(f'Sheet: {sheet.title}\n' + '\n'.join(rows))
    body = '\n\n'.join(out)
    if len(body.strip()) < MIN_TEXT:
        raise ImportProblem('This Excel file has almost nothing in it. Upload the sheet with your quotation.')
    return body[:MAX_TEXT]


# ── 2. Read it ───────────────────────────────────────────────────────────────

class Line(BaseModel):
    name: str
    description: str
    size: str
    unit: str
    quantity: Optional[float]
    rate: Optional[float]
    amount: Optional[float]


class Section(BaseModel):
    name: str
    lines: List[Line]


class Charge(BaseModel):
    name: str
    amount: float


class Quote(BaseModel):
    job_title: str
    sections: List[Section]
    charges: List[Charge]
    subtotal: Optional[float]
    discount: Optional[float]
    gst_percent: Optional[float]
    grand_total: Optional[float]
    payment_terms: str
    notes_and_terms: str


PROMPT = """This is the text of a quotation (BOQ) that a business (interiors, furniture, kitchens, windows, glass work…)
sent to a customer, taken from a PDF or an Excel sheet, so table columns may be run together.
Read it into the structure:
- sections: the rooms or section headings, in order, each with its priced lines.
- each line: the item's name as written, without its serial number ("TV Unit Base Drawers", not "1.1 TV Unit Base
  Drawers"; keep the words that tell similar items apart); the specification or description as written ('' if none);
  the size as written ('' if none); the unit as written (sqft, rft, nos, LS…); quantity = what the rate is
  multiplied by to give the amount (the area or length when the rate is per sqft or rft, else the count); rate and
  amount exactly as printed.
- charges: fees or extra amounts added after the lines, such as a design fee, contingency or transport. Not taxes.
- subtotal (total of the lines, before discount, charges and tax), discount, gst_percent, grand_total: as printed.
- job_title: what the job is, in 2–4 words ("3 BHK interiors", "Pharmacy interiors", "Kitchen and wardrobes").
- payment_terms and notes_and_terms: their own wording, '' if none.
Copy numbers exactly (Indian format: 1,45,000 = 145000). Never estimate or invent a number: use null when one isn't
printed. Leave out the customer's name, phone number and address."""


def read(body):
    """Claude's reading of the quotation text → (Quote, usage). Raises ImportProblem."""
    import anthropic
    client = anthropic.Anthropic(api_key=api_key(), timeout=TIMEOUT, max_retries=1)
    try:
        r = client.messages.parse(model=MODEL, max_tokens=16000, output_format=Quote,
                                  messages=[{'role': 'user', 'content': f'{PROMPT}\n\n<quotation>\n{body}\n</quotation>'}])
    except anthropic.APIStatusError as exc:
        raise ImportProblem('Reading the quotation failed just now. Try again in a minute.') from exc
    except anthropic.APIConnectionError as exc:
        raise ImportProblem('Reading the quotation took too long or couldn\'t connect. Try again in a minute.') from exc
    if r.stop_reason == 'max_tokens':
        raise ImportProblem('This quotation is too long to read in one go. Upload a shorter one (one job).')
    if r.stop_reason == 'refusal' or r.parsed_output is None:
        raise ImportProblem("This file couldn't be read as a quotation.")
    return r.parsed_output, {'in': r.usage.input_tokens, 'out': r.usage.output_tokens}


def read_upload(filename, raw):
    """The whole of steps 1–2 for an upload, within the limit: the review data (see tidy). Raises ImportProblem."""
    if left()[0] <= 0:
        raise ImportProblem(f'You\'ve used all your quotation imports ({left()[1]}). Message us on WhatsApp if you '
                            'need another.')
    kind, body = file_text(filename, raw)
    try:
        quote, used = read(body)
    except ImportProblem as exc:
        _log(kind, False, note=str(exc))
        raise
    data = tidy(quote)
    lines = sum(len(s['lines']) for s in data['sections'])
    if not lines:
        _log(kind, False, tokens_in=used['in'], tokens_out=used['out'], note='no lines')
        raise ImportProblem("No priced lines were found in this file. Is it a quotation with items and amounts?")
    _count()
    _log(kind, True, lines, used['in'], used['out'])
    return data


# ── 3. Review ────────────────────────────────────────────────────────────────

# "1.1 ", "3. ", "4) ", "a) ", "(iv) " in front of a name (a bare "2 " is part of it: "2 BHK wardrobe")
_NUMBERING = re.compile(r'^\s*(?:\(?\d+(?:\.\d+)+[.)]?|\(?\d+[.)]|\(?[A-Za-z][.)]|\(?[ivxIVX]+[.)])\s+')


def _name(value):
    name = _NUMBERING.sub('', ' '.join(str(value or '').split()))
    return (name[:1].upper() + name[1:])[:100]


def _section(value):
    """A section heading: "A MASTER BEDROOM" → "Master Bedroom"."""
    name = _name(re.sub(r'^\s*[A-Z]\s+(?=[A-Z])', '', ' '.join(str(value or '').split())))
    return name.title() if name.isupper() else name


def _num(value):
    """A number (whole ones as int, so 72 shows as 72, not 72.0), or None."""
    try:
        n = round(float(value), 2) if value is not None and value != '' else None
    except (TypeError, ValueError):
        return None
    return int(n) if n is not None and n == int(n) else n


def tidy(quote):
    """The reading as plain data for the review page: names without serial numbers, units in the app's words, a
    rate worked out from amount ÷ quantity when only the amount is printed, the size kept in the description."""
    from . import boq
    sections = []
    for s in quote.sections:
        lines = []
        for l in s.lines:
            name = _name(l.name)
            if not name:
                continue
            qty, rate, amount = _num(l.quantity), _num(l.rate), _num(l.amount)
            if rate is None and amount is not None:
                rate, qty = (_num(amount / qty), qty) if qty else (amount, 1)
            elif rate and amount is not None and (qty is None or abs(qty * rate - amount) > 1):
                qty = _num(round(amount / rate, 3))   # the printed amount wins: area × count × rate, say
            desc = ' '.join((l.description or '').split())
            if desc.lower() == name.lower():
                desc = ''
            if l.size.strip():
                desc = (desc + ' · ' if desc else '') + 'Size: ' + ' '.join(l.size.split())
            lines.append({'name': name, 'description': desc[:500], 'unit': boq.clean_unit(l.unit),
                          'quantity': qty, 'rate': rate, 'amount': amount,
                          'no_amount': amount is None and rate is not None})   # maybe not in their total
        if lines:
            sections.append({'name': _section(s.name) or 'Items', 'lines': lines})
    charges = [{'name': _name(c.name), 'amount': _num(c.amount)} for c in quote.charges if _name(c.name) and _num(c.amount)]
    return {'job_title': ' '.join(quote.job_title.split())[:60] or 'My quotation', 'sections': sections,
            'charges': charges, 'subtotal': _num(quote.subtotal), 'discount': _num(quote.discount) or 0,
            'gst_percent': _num(quote.gst_percent), 'grand_total': _num(quote.grand_total),
            'payment_terms': quote.payment_terms.strip()[:3000], 'notes': quote.notes_and_terms.strip()[:5000]}


def from_form(form):
    """The review page's form → the same plain data (lines left unticked are dropped)."""
    sections = []
    for i in range(min(int(form.get('sections') or 0), 60)):
        lines = []
        for j in range(min(int(form.get(f's{i}_lines') or 0), 200)):
            p = f's{i}_{j}_'
            if not form.get(p + 'keep') or not _name(form.get(p + 'name')):
                continue
            lines.append({'name': _name(form.get(p + 'name')), 'description': (form.get(p + 'description') or '').strip()[:500],
                          'unit': (form.get(p + 'unit') or '').strip(), 'quantity': _num(form.get(p + 'quantity')),
                          'rate': _num(form.get(p + 'rate'))})
        if lines:
            sections.append({'name': _section(form.get(f's{i}_name')) or 'Items', 'lines': lines})
    charges = []
    for k in range(min(int(form.get('charges') or 0), 20)):
        if form.get(f'c{k}_keep') and _name(form.get(f'c{k}_name')) and _num(form.get(f'c{k}_amount')):
            charges.append({'name': _name(form.get(f'c{k}_name')), 'amount': _num(form.get(f'c{k}_amount'))})
    return {'job_title': ' '.join((form.get('job_title') or '').split())[:60] or 'My quotation', 'sections': sections,
            'charges': charges, 'discount': _num(form.get('discount')) or 0, 'gst': _num(form.get('gst')),
            'payment_terms': (form.get('payment_terms') or '').strip()[:3000], 'notes': (form.get('notes') or '').strip()[:5000]}


# ── 4. Set up ────────────────────────────────────────────────────────────────

def price_list(data):
    """The reviewed data as a price list file (app/price_list.py): saved lines (one per item name; the same name at
    another rate gets its section's name added), a starting quotation with those sections, the quotation rebuilt,
    payment terms and notes, and GST if it was on."""
    from . import boq, quote_gst
    lines, by_name = [], {}

    def line_for(l, section):
        rate = l['rate'] or 0
        name, n = l['name'], 1
        while name.lower() in by_name and by_name[name.lower()]['rate'] != rate:     # same name, another rate
            n += 1
            name = (f"{l['name']} ({section})" if n == 2 else f"{l['name']} ({section} {n - 1})")[:100]
        key = name.lower()
        if key not in by_name:
            by_name[key] = {'name': name, 'description': l['description'], 'unit': boq.clean_unit(l['unit']),
                            'rate': rate, 'group': section[:60]}
            lines.append(by_name[key])
        return by_name[key]['name']

    rooms, quote_rooms = [], []
    sections = data['sections'] + ([{'name': 'Other charges', 'lines': [
        {'name': c['name'], 'description': '', 'unit': 'lump sum', 'quantity': 1, 'rate': c['amount']}
        for c in data['charges']]}] if data['charges'] else [])
    for s in sections:
        names, items = [], []
        for l in s['lines']:
            name = line_for(l, s['name'])
            if name not in names:
                names.append(name)
            items.append({'name': name, 'description': l['description'], 'unit': boq.clean_unit(l['unit']),
                          'quantity': l['quantity'] or 0, 'rate': l['rate'] or 0})
        rooms.append({'name': s['name'], 'items': names})
        quote_rooms.append({'name': s['name'], 'items': items})
    out = {'pricing': 'item', 'lines': lines,
           'templates': [{'name': data['job_title'], 'rooms': rooms}],
           'quotations': [{'customer_name': f"{data['job_title']} (your old quotation)", 'discount': data['discount'],
                           'rooms': quote_rooms}],
           'settings': {k: v for k, v in (('payment_terms', data['payment_terms']), ('important_notes', data['notes'])) if v}}
    if data.get('gst') in quote_gst.RATES:
        out['options'] = {'gst': data['gst']}
    return out


def set_up(data):
    """Load the reviewed quotation into this business (its saved lines, template, settings, the quotation).
    Returns (summary, the rebuilt quotation's id or None). Raises price_list.PriceListError."""
    from . import price_list as PL
    from .models import Project
    result = PL.apply(PL.parse(json.dumps(price_list(data)).encode()), remove_starter=True, replace_settings=False)
    made = None
    if result['quotations']:
        name = result['quotations'][0]
        made = next((p.id for p in Project.query.order_by(Project.id.desc()).limit(20) if p.customer_name == name), None)
    return result, made
